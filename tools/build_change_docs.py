#!/usr/bin/env python3
"""Build the technical and nontechnical Serv Pi 2.1 change reports."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
DATE = "August 5, 2026"
BLUE = "174A7E"
MID_BLUE = "2F6DA1"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F4F8FB"
GREEN = "2E7D32"
LIGHT_GREEN = "EAF5EC"
AMBER = "9A6700"
LIGHT_AMBER = "FFF4CE"
RED = "A61B1B"
LIGHT_RED = "FDECEC"
INK = "1F2933"
GRAY = "5F6B76"
LIGHT_GRAY = "EEF1F4"
WHITE = "FFFFFF"


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, amount in (("top", top), ("start", start),
                         ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell, dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def _cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:cantSplit")
    tr_pr.append(tag)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_repeat_table_header(table) -> None:
    if table.rows:
        _repeat_header(table.rows[0])


def _page_number(paragraph) -> None:
    paragraph.add_run("SERV  •  ")
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _setup(document: Document, short_title: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.36)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in (
        ("Title", 24, BLUE, 0, 10),
        ("Subtitle", 11, GRAY, 0, 8),
        ("Heading 1", 16, BLUE, 16, 7),
        ("Heading 2", 13, MID_BLUE, 11, 5),
        ("Heading 3", 11.5, BLUE, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = f"SERV  /  {short_title.upper()}"
    hp.style = styles["Caption"]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _page_number(fp)
    for run in fp.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    props = document.core_properties
    props.author = "Serv Engineering"
    props.company = "Serv"
    props.subject = "Serv Pi release 2.1.0"
    props.keywords = "Raspberry Pi, ALPR, reliability, optimization, field camera"


def _masthead(document: Document, title: str, subtitle: str, audience: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left, right = table.rows[0].cells
    _set_cell_width(left, 6800)
    _set_cell_width(right, 2560)
    _shade(left, BLUE)
    _shade(right, LIGHT_BLUE)
    for cell in (left, right):
        _cell_margins(cell, top=150, bottom=150, start=180, end=180)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("SERV")
    r.bold = True
    r.font.size = Pt(19)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p2 = left.add_paragraph("FIELD VISION SYSTEMS")
    p2.paragraph_format.space_after = Pt(0)
    p2.runs[0].font.size = Pt(8)
    p2.runs[0].font.bold = True
    p2.runs[0].font.color.rgb = RGBColor.from_string("C9DDF0")

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = rp.add_run("RELEASE 2.1.0")
    rr.bold = True
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor.from_string(BLUE)
    rp2 = right.add_paragraph(DATE)
    rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp2.runs[0].font.size = Pt(8.5)
    rp2.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

    document.add_paragraph(title, style="Title")
    document.add_paragraph(subtitle, style="Subtitle")
    meta = document.add_table(rows=2, cols=2)
    meta.autofit = False
    widths = [1900, 7460]
    for row, (label, value) in zip(meta.rows, (
        ("Audience", audience),
        ("Release position", "Short 2.1.0 field canary passed health and was rolled back to 2.0.0 pending cooling, lane calibration, and labeled-video validation"),
    )):
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[index])
            _cell_margins(cell, top=65, bottom=65)
        _shade(row.cells[0], LIGHT_GRAY)
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)


def _callout(document: Document, title: str, text: str,
             fill: str = LIGHT_BLUE, accent: str = BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    _cant_split(table.rows[0])
    cell = table.cell(0, 0)
    _shade(cell, fill)
    _cell_margins(cell, top=130, bottom=130, start=170, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        run.font.size = Pt(10)


def _bullets(document: Document, items: list[str], level: int = 0) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True


def _numbered(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Number")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True


def _table(document: Document, headers: list[str], rows: list[list[str]],
           widths: list[int], small: bool = False) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        _set_cell_width(cell, widths[index])
        _shade(cell, BLUE)
        _cell_margins(cell)
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(8.5 if small else 9)
    _set_repeat_table_header(table)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        _cant_split(row)
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            _set_cell_width(cell, widths[index])
            _cell_margins(cell)
            if row_index % 2:
                _shade(cell, PALE_BLUE)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8.4 if small else 9)


def _status_strip(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    items = [
        ("CODE", "Complete", LIGHT_GREEN, GREEN),
        ("DESKTOP TESTS", "218 passed", LIGHT_GREEN, GREEN),
        ("FIELD CANARY", "Short pass / hold", LIGHT_AMBER, AMBER),
    ]
    for cell, (label, value, fill, color) in zip(table.rows[0].cells, items):
        _set_cell_width(cell, 3120)
        _shade(cell, fill)
        _cell_margins(cell, top=105, bottom=105)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(color)
        p2 = cell.add_paragraph(value)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        for rr in p2.runs:
            rr.font.size = Pt(10.5)
            rr.bold = True
            rr.font.color.rgb = RGBColor.from_string(color)


def build_technical() -> Path:
    document = Document()
    _setup(document, "Technical Change Report")
    _masthead(
        document,
        "Serv Pi 2.1 — Technical Change Report",
        "Logic audit, Raspberry Pi optimization, reliability hardening, and field-validation plan",
        "Engineering, operations, backend, and field deployment owners",
    )
    _status_strip(document)

    document.add_heading("1. Executive engineering outcome", level=1)
    document.add_paragraph(
        "The codebase has been converted from a promising camera prototype into a safer "
        "release candidate. The largest correctness risks were not raw model speed; they "
        "were duplicate lifecycle paths, same-frame votes masquerading as independent "
        "evidence, ambiguous vehicle association, a delivery queue that could block behind "
        "bad work, and health checks that could report success before a real frame passed "
        "through inference. Those paths are now explicit, bounded, and testable."
    )
    _callout(
        document,
        "Release decision",
        "The short code-only canary passed its real camera/inference health gate and held "
        "5.01 fps, but the unit has no cooler and operates near the Pi 5 thermal-control "
        "boundary. Keep 2.0.0 active until cooling is installed and the labeled video can "
        "be transferred over Wi-Fi or USB. Do not claim 98% or deploy fleet-wide.",
        LIGHT_AMBER,
        AMBER,
    )

    document.add_heading("2. Scope and operating contract", level=1)
    document.add_paragraph(
        "This review covered capture, inference adapters, normalization and hashing, "
        "multi-frame voting, vehicle presence, visit/event delivery, telemetry, systemd "
        "hardening, cellular checks, installation/update scripts, the vehicle-profile "
        "prototype, tests, deployment guidance, a read-only audit, and a supervised code-only "
        "field canary. The Pi was restored to 2.0.0 afterward; the 2.1.0 package remains on-device "
        "as a recoverable candidate. No dependency or systemd-unit upgrade was performed."
    )
    _table(document, ["Production contract", "2.1 behavior"], [
        ["Vehicle counting", "The presence tracker is the single production authority for a vehicle lifecycle and drive-off."],
        ["Plate identity", "A normalized string must agree across at least two independent captures; the template uses three."],
        ["Plate privacy", "Only a validated 64-character hash leaves the identity boundary. Raw strings are not written at INFO."],
        ["Visit delivery", "SQLite-backed, priority-aware, idempotency-tagged, retrying, and dead-lettered when permanently invalid."],
        ["Health", "A release becomes healthy only after real camera capture and successful inference complete repeatedly."],
        ["Vehicle profile", "Color stays enabled in the measured production profile. Make/model remains off until its frame cost is acceptable and the cloud schema is approved."],
    ], [2500, 6860])

    document.add_heading("3. Correctness and logic changes", level=1)
    _table(document, ["Risk found", "Change made", "Why it matters"], [
        ["Heartbeat thread overrode Python Thread._stop", "Renamed the event and made shutdown/join safe.", "Prevents an orderly service stop from failing or hanging."],
        ["Multiple detections in one SDK response could satisfy a vote", "Collapse identical normalized candidates per capture; keep only the strongest observation.", "A multi-frame identity now requires independent camera moments."],
        ["Dwell time survived missing frames", "Reset continuous dwell when the candidate is absent.", "A vehicle cannot earn dwell credit while not observed."],
        ["Degenerate plate boxes failed zone checks", "Use the vehicle box when a plate box is absent or zero-sized.", "Plateless vehicles remain countable within the calibrated lane."],
        ["Two unidentified-vehicle paths could disagree or double count", "Production disables the detector-side legacy path; presence owns the lifecycle.", "One car produces at most one unidentified outcome."],
        ["Greedy matching reused an observation/track", "Use one-to-one association and collapse near-identical SDK boxes.", "Prevents duplicate records and identity swaps within a capture."],
        ["Plate attribution could guess between vehicles", "Attach by exact normalized identity, strong spatial match, or a sole unambiguous unresolved track; otherwise refuse.", "It is safer to emit unknown than assign a customer to the wrong car."],
        ["Make and model could be chosen independently", "Treat make/model as one confidence-scored pair.", "Prevents creation of a vehicle combination the classifier never proposed."],
        ["Stale field config could report an old release", "Installed package version is authoritative; a mismatch is logged.", "Cloud diagnostics show the software that is actually running."],
    ], [2080, 3590, 3690], small=True)

    document.add_heading("4. Delivery resilience and API behavior", level=1)
    _bullets(document, [
        "The SQLite queue uses write-ahead logging and normal synchronous mode for practical durability without forcing a full filesystem sync per event.",
        "Rows carry priority, next-attempt time, attempt count, and a stable idempotency key. Visits outrank detections; detections outrank heartbeats.",
        "Retries use bounded exponential backoff with jitter, and ready work can proceed around a delayed row. This removes head-of-line blocking.",
        "Clearly permanent request errors are retained in a bounded dead-letter table. Authentication failures remain retryable and fail the startup preflight instead of discarding future visits.",
        "When the queue is full, low-priority telemetry cannot evict a visit. Queue depth, dead letters, oldest age, and maximum attempts are reported in heartbeat telemetry.",
    ])
    _callout(
        document,
        "Backend confirmation required",
        "The client now sends an Idempotency-Key header. The API owner should confirm the "
        "server stores and honors it for the intended retention window. Client-side use is "
        "safe, but duplicate prevention is only end-to-end when the server enforces the key.",
    )

    document.add_heading("5. Raspberry Pi performance and stability", level=1)
    _table(document, ["Area", "Implementation", "Expected effect / measurement"], [
        ["Camera path", "Picamera2 video configuration, RGB888 frames, fixed FrameDurationLimits, four buffers, and queue disabled.", "Reduces conversion and stale-frame risk; confirm on the installed sensor."],
        ["Loop pacing", "Monotonic/performance clocks pace to capture.target_fps; legacy capture_interval_ms is ignored.", "A single 5 fps budget equals 200 ms/frame instead of two competing delays."],
        ["Memory", "Removed forced garbage collection from the hot loop.", "Avoids periodic pause spikes; bounded structures control growth."],
        ["CPU scheduling", "Service-level Nice=-5 replaces ineffective per-user adjustment.", "Makes scheduling policy visible and repeatable under systemd."],
        ["Classifier profile", "Vehicle box and color on; LPCI, VMMR, and VBSR off by default.", "Keeps the previously measured profile near 106 ms/frame; heavy profile was about 268 ms/frame and missed the 5 fps budget."],
        ["Dependencies", "Base install excludes FastALPR/ONNX unless that engine is selected.", "Less disk, install time, and dependency exposure on UltimateALPR units."],
        ["Telemetry", "Stable process sampler; CPU, RSS, threads, swap, temperature, throttling, clock, disk writes, queue state, and p50/p95 timings.", "Field bottlenecks become measurable instead of inferred."],
        ["Failure recovery", "Repeated capture/inference failures terminate the process so systemd can restart it; contract/license errors are immediately fatal.", "Avoids a green process that is no longer processing real images."],
    ], [1600, 4420, 3340], small=True)

    document.add_heading("6. Security and privacy hardening", level=1)
    _bullets(document, [
        "API keys must be exactly 64 lowercase hexadecimal characters, and API endpoints must use HTTPS without embedded credentials, query strings, or fragments.",
        "License verification is fail-closed. Trial-mode plate text containing a redaction marker is treated as fatal before normalization or hashing.",
        "The preview server binds only to 127.0.0.1 and is intended to be reached through an SSH tunnel; it cannot be opened to the LAN by an environment variable.",
        "Vehicle-record captures and temporary configuration files use owner-only permissions. Raw replay/record files are ignored by Git.",
        "Broad modem sudo access was replaced with a fixed-function telemetry wrapper. systemd now adds kernel, namespace, SUID, personality, and private-temp restrictions.",
        "Operational logs no longer print normalized plate text at INFO. Suppression diagnostics use a short hash prefix.",
    ])
    _callout(
        document,
        "Open privacy decision",
        "The existing identity remains plain SHA-256 of the normalized plate. This preserves "
        "cloud compatibility but permits offline guessing of a small plate namespace if hashes "
        "leak. Moving to keyed HMAC is recommended as a coordinated backend migration—not as "
        "an isolated Pi change—because it changes every stored identity and matching workflow.",
        LIGHT_RED,
        RED,
    )

    document.add_heading("7. Installer, updater, and rollback", level=1)
    _bullets(document, [
        "UltimateALPR SDK checkout is pinned to commit 972506c84ad286bb64331943682084a078f43494 and verified after checkout.",
        "Install/update stages packages before swapping them into place, preserves the prior package, and restores it when validation or health checks fail.",
        "A staged config is never installed unless explicitly requested; existing per-device credentials and calibration remain in place by default.",
        "The update health sentinel requires successful real capture and inference. Log screening treats configuration, authentication, detector, license, and camera failures as fatal.",
        "Manual staging deployments create an override marker so the pull updater cannot immediately erase an emergency hotfix.",
        "A read-only field audit script gathers hardware, thermal, camera, service, SDK, queue, network, and selected health information without printing credentials or queue payloads.",
    ])
    document.add_paragraph(
        "Rollback boundary: code and configuration are recoverable through the package backup. "
        "A pip dependency upgrade is not automatically uninstalled by code rollback. Perform the "
        "first canary with a disk image or package snapshot available if dependency changes are included."
    )

    document.add_heading("8. Vehicle-profile prototype status", level=1)
    document.add_paragraph(
        "The separate vehicle_records module is substantially hardened but remains a prototype "
        "until its payload contract is approved. It now generates boot/session-scoped record IDs, "
        "uses one-to-one tracking, refuses ambiguous plate attachment, keeps color by strongest "
        "evidence, updates make/model atomically, freezes records at send time, distinguishes "
        "confirmed / low-confidence / no-plate-visible outcomes, and emits plateless records by default."
    )
    _callout(
        document,
        "Integration boundary",
        "Production continues to send the established /visit and /detection-event shapes. Do not "
        "wire the new vehicle profile to the cloud until backend ownership, schema versioning, "
        "retention, and reconciliation behavior are explicitly approved.",
    )

    document.add_heading("9. Verification evidence", level=1)
    _table(document, ["Suite", "Result", "Coverage emphasis"], [
        ["Board trigger", "12 / 12", "Settle, fallback, slot ownership, no re-fire"],
        ["Dwell suppression", "14 / 14", "Presence, absence, return, confidence extraction"],
        ["Dwell end-to-end", "4 / 4", "Repeated reads and return after departure"],
        ["Motion association", "8 / 8", "Fast approach, two cars, no swaps, jitter"],
        ["Occlusion hold", "16 / 16", "FIFO hold, reappearance, bounded queue advance"],
        ["Pipeline features", "76 / 76", "ROI, gates, dedup, SDK contract, configuration"],
        ["Release hardening", "16 / 16", "Thread stop, queue, config merge, boxes, hasher"],
        ["Track report", "7 / 7", "Fragmentation/near-miss instrumentation"],
        ["Vehicle records", "65 / 65", "Invariants, association, freezing, wire shape"],
        ["Python + shell validation", "Pass", "compileall and bash -n across all scripts"],
        ["Field startup health", "Pass", "2.1.0 reached real camera + inference health; license probe was unmasked"],
        ["Short Pi throughput", "5.01 fps", "301 frames in 60 seconds at a 5.00 fps target"],
        ["Queue recovery", "321 sent / 13 retained", "Valid backlog drained; permanent legacy HTTP 400 rows were preserved in dead letter"],
        ["OpenCV motion boxes", "Lane run pending", "Indoor camera placement cannot validate lane geometry or traffic counting"],
    ], [2500, 1450, 5410], small=True)
    document.add_paragraph(
        "Total: 218 assertions passed and zero failed. These deterministic tests validate logic; "
        "they do not substitute for camera geometry, lighting, thermals, cellular loss, or real SDK timing."
    )
    _callout(
        document,
        "Field-canary snapshot — serv-ljs-001, August 5",
        "Raspberry Pi 5 (4 GiB) with IMX296 camera. Before deployment, the long-running 2.0.0 "
        "service measured 81.2 °C. During the 50-second 2.1.0 guard window, temperature ranged "
        "78.5–80.7 °C with no current throttle bits; one process snapshot was 69.0% CPU and "
        "approximately 282 MiB RSS. After rollback, a comparable 50-second 2.0.0 sample ranged "
        "77.9–79.6 °C at 69.1–69.3% CPU and approximately 276 MiB RSS. The run is too short to "
        "claim a thermal improvement. Historical frequency-cap, throttling, and soft-temperature "
        "flags remain recorded since boot (get_throttled=0xe0000).",
        LIGHT_AMBER,
        AMBER,
    )
    document.add_paragraph(
        "Queue finding: 2.1.0 started with 334 queued rows. It delivered 321 and moved 13 "
        "legacy 2.0.0 /detection-event rows returning HTTP 400 into recoverable dead-letter "
        "storage. One poison row had accumulated 48,319 attempts. This confirms the new "
        "head-of-line and permanent-error handling, while exposing a backend contract decision "
        "for vehicle_only and group_confidence_floor outcomes."
    )

    document.add_heading("10. Field canary acceptance plan", level=1)
    _numbered(document, [
        "Install a Pi 5 Active Cooler or a thermally equivalent enclosure solution before another sustained inference run.",
        "Connect the Pi to a non-cellular Wi-Fi or Ethernet route, or attach USB storage. The 2.7 GB video must not be sent over the current cellular route.",
        "Run the labeled 30-event clip against 2.0.0 for a baseline. The private ground truth contains 21 conventional plate-visible events, 8 no-plate events, and 1 nonstandard placard.",
        "Redeploy the preserved 2.1.0 candidate with the automatic health gate and rollback backup, then rerun the identical clip and compare exact-match accuracy, event count, latency, CPU, RSS, temperature, clock, and throttle flags.",
        "When the camera is installed at a drive-through, review a privacy-safe preview and calibrate the lane, plate, and menu-board zones before enabling lifecycle acceptance tests.",
        "Run the complete live pipeline—not only model inference—for at least 30 minutes at the configured 5 fps after cooling and zone calibration.",
        "Review false matches, missed vehicles, duplicate visits, ambiguous associations, queue growth, p95 latency, thermals, and throttling before expanding beyond the canary.",
    ])
    _table(document, ["Gate", "Canary acceptance criterion"], [
        ["Throughput", "Sustained configured 5 fps; p95 end-to-end loop no greater than 200 ms, or an explicitly approved lower target if the sensor/SDK cannot maintain it."],
        ["Thermals/power", "No current or newly latched under-voltage/thermal-throttle bits; target sustained temperature <75 °C, preferably <70 °C for environmental margin."],
        ["Vehicle recall", "Human lane count reconciles to system vehicle lifecycles; every miss is explainable and the measured recall meets the agreed product target."],
        ["Plate precision", "Zero knowingly wrong hashes or cross-vehicle assignments. Unknown is preferable to wrong."],
        ["98% claim", "Use at least 200 eligible, plate-visible encounters across representative day/night/weather conditions before presenting the result as credible; report the sample size and confidence interval."],
        ["Duplicates", "No second visit for a continuously present vehicle; returning vehicles count after a real departure."],
        ["Network outage", "Visits remain queued during loss, heartbeats cannot evict them, and the backlog drains without duplicate cloud effects after recovery."],
        ["Recovery", "Camera/inference failure causes a visible unhealthy state and systemd restart; failed deployment rolls back to the prior package."],
    ], [2200, 7160], small=True)

    document.add_page_break()
    document.add_heading("11. Residual risks and required decisions", level=1)
    _table(document, ["Item", "Owner / next decision", "Release impact"], [
        ["Sustained camera/Pi performance", "Field owner installs cooling; engineering runs 30-minute live and labeled-video benchmarks.", "Short canary passed, but sustained validation blocks rollout."],
        ["Rejected detection outcomes", "Backend decides whether vehicle_only and group_confidence_floor are valid API outcomes.", "13 legacy rows remain preserved; do not replay until the contract is resolved."],
        ["HMAC identity migration", "Backend + privacy decide key ownership and data migration.", "Plain SHA remains compatible but is the main privacy debt."],
        ["Idempotency enforcement", "API owner confirms storage/retention of Idempotency-Key.", "Needed for end-to-end duplicate prevention."],
        ["Vehicle-profile payload", "Product/backend approve schema, retention, null semantics, and send trigger.", "Color/make/model profile stays disconnected until approved."],
        ["Make/model performance", "Field benchmark with VMMR on only if product value justifies lower fps.", "Default remains off."],
        ["Dependency rollback", "Operations chooses image/package snapshot practice for canary.", "Code rollback alone may leave upgraded Python packages."],
        ["Camera zones", "Field owner provides a safe preview and confirms lane/board bounds.", "Uncalibrated geometry invalidates counting claims."],
    ], [2200, 4110, 3050], small=True)
    _callout(
        document,
        "Current operational state",
        "serv-ljs-001 is running 2.0.0 with an empty ready queue. The 13 rejected legacy "
        "records remain recoverable in dead-letter storage, the 2.1.0 candidate is preserved "
        "on the Pi, and automatic updates are intentionally paused by the manual-override "
        "marker until the next supervised acceptance run.",
        LIGHT_AMBER,
        AMBER,
    )
    document.add_heading("Evidence retained", level=2)
    _bullets(document, [
        "FIELD_AUDIT_serv-ljs-001_2026-08-04.md — read-only hardware, thermal, service, SDK, queue, and network baseline.",
        "FIELD_CANARY_serv-ljs-001_2026-08-05.md — exact deployment, rollback, runtime, queue, and final-state record.",
        "dtfootagefinal_private_ground_truth.xlsx — private 30-event timeline and conventional plate labels; excluded from Git.",
    ])

    path = OUT / "Serv_Pi_Technical_Change_Report_2026-08-05.docx"
    document.save(path)
    return path


def build_nontechnical() -> Path:
    document = Document()
    _setup(document, "Change Summary")
    _masthead(
        document,
        "Serv Pi 2.1 — Change Summary",
        "What changed, why it matters, what is proven, and what is needed next",
        "Founders, operations, product, and field partners",
    )
    _status_strip(document)

    document.add_heading("The short version", level=1)
    document.add_paragraph(
        "The camera software is now safer, more predictable, and easier to measure. It is "
        "better protected against counting one car twice, assigning a plate to the wrong car, "
        "losing visits when the internet is down, reporting itself healthy when the camera is "
        "not actually working, and exposing sensitive information in logs or previews."
    )
    _callout(
        document,
        "Where we are today",
        "The code work is complete, 218 automated checks passed, and a short supervised test ran "
        "on Jasper's Pi. Version 2.1 reached healthy camera + recognition operation and held its "
        "5 fps target. We restored 2.0 afterward because the Pi has no cooler, the camera is "
        "indoors rather than aimed at a lane, and its Wi-Fi is disconnected.",
        LIGHT_AMBER,
        AMBER,
    )

    document.add_heading("What changed", level=1)
    _table(document, ["Change", "Why it matters to the business"], [
        ["One system now owns each car's journey through the lane", "Reduces double-counts and conflicting answers about whether a car was seen."],
        ["A plate must agree across multiple camera moments", "A single bad frame cannot become a customer identity."],
        ["The system refuses uncertain car-to-plate matches", "It will report unknown instead of confidently attaching a customer to the wrong vehicle."],
        ["Visits are protected during internet outages", "Important visit data is kept ahead of routine health messages and retried safely."],
        ["Health now means the camera and recognition actually worked", "Operations can distinguish a running process from a functioning product."],
        ["The Pi workload is simpler and measurable", "Camera pacing, temperature, power warnings, memory, CPU, and real processing time are visible."],
        ["Updates can roll back when the camera does not recover", "A failed release is less likely to strand a unit in the field."],
        ["Remote viewing and logs are more private", "The preview stays on the Pi unless reached through an SSH tunnel, and plate text is kept out of normal logs."],
    ], [3180, 6180])

    document.add_heading("How this addresses the planning checklist", level=1)
    _table(document, ["Planning goal", "Current answer"], [
        ["At least 98% plate recognition", "The software has stronger voting and fewer false confirmations. The percentage is not yet certified; it must be measured on the field camera with labeled traffic."],
        ["Recognize every car that goes through", "The system now counts vehicle lifecycles separately from plate success, including a car whose plate is missing or unreadable. Real lane recall still requires a human comparison trial."],
        ["Car profile with plate hash, color, make, model", "A hardened profile prototype exists. Color fits the current speed budget. Make/model is kept off by default because the measured heavy profile was too slow. The new profile is not sent to the cloud until the schema is approved."],
        ["Dispose of cars that leave without dwelling", "The presence logic tracks board arrival, settling, departure, and drive-off. The exact board/lane zone must be calibrated on each camera."],
        ["Avoid background cars", "The design uses a closed lane zone, vehicle evidence, dwell/board behavior, and movement. Camera angle and zone calibration remain essential."],
        ["Correct a weak character on later frames", "Multiple frames now vote on the exact normalized plate. Same-frame duplicates do not get extra votes."],
        ["Use Pi resources efficiently", "The default recognition profile stays inside the known frame budget, unnecessary dependencies are skipped, and field telemetry will reveal the real bottleneck."],
    ], [3180, 6180], small=True)

    document.add_heading("What is proven", level=1)
    _bullets(document, [
        "218 automated logic checks passed with zero failures.",
        "All Python files compile, and all field shell scripts pass syntax validation.",
        "The tests cover repeated plate reads, return visits, two cars close together, occlusion, queue outages, permanent API errors, bad configuration, ambiguous plate assignment, and record privacy invariants.",
        "On the real Pi 5 and IMX296 camera, 2.1 became healthy, reported the correct version, and processed 301 frames in 60 seconds (5.01 fps).",
        "The new queue logic delivered 321 valid waiting messages instead of remaining blocked behind one bad message. It preserved 13 permanently rejected old messages for review rather than deleting them.",
        "The supplied recording now has private ground truth for all 21 conventional plate-visible events; the 30-event timeline also includes 8 no-plate vehicles and 1 nonstandard placard.",
    ])

    document.add_heading("What is not proven yet", level=1)
    _bullets(document, [
        "Sustained speed and temperature after cooling. The short test stayed near 79–81 °C and is not enough to clear the unit for long production runs.",
        "The 98% recognition claim. A small successful demo is not enough; use at least 200 eligible plate-visible encounters and publish the sample size.",
        "That every lane vehicle is counted. A human must compare the real traffic count against the system count.",
        "The exact backend rule for two old diagnostic outcomes that are currently rejected with HTTP 400.",
        "Lane and menu-board behavior. Jasper's camera is currently indoors, so those zones cannot be calibrated or judged yet.",
    ])
    _callout(
        document,
        "Important expectation",
        "The safer system sometimes chooses “unknown” instead of guessing. That may look like "
        "lower recognition in a small demo, but it prevents the more damaging error: assigning "
        "an order or customer to the wrong vehicle.",
        LIGHT_AMBER,
        AMBER,
    )

    document.add_heading("Recommended next step", level=1)
    _numbered(document, [
        "Install a Raspberry Pi 5 Active Cooler, or an enclosure cooling solution with equivalent performance.",
        "Connect Jasper's Pi to non-cellular Wi-Fi/Ethernet, or attach a USB drive containing the 2.7 GB validation video.",
        "Run the same 30-event video through 2.0 and 2.1 while measuring exact plate matches, vehicle count, CPU, memory, temperature, clock speed, and throttling.",
        "Resolve the backend rule for the 13 preserved old messages before anyone replays them.",
        "Later, when the camera is at a real drive-through, review the view and set the lane and menu-board boundaries.",
        "Build the accuracy set toward at least 200 eligible plate-visible encounters before making a 98% claim.",
    ])

    document.add_heading("What we need from Carlos and Theo", level=1)
    _table(document, ["Needed", "Purpose"], [
        ["Cooling installed", "Heat is the main hardware blocker. Tell engineering which cooler/enclosure is installed and when the Pi is ready for a 30-minute load."],
        ["Non-cellular file path", "Ask Jasper to connect the Pi to Wi-Fi/Ethernet or put the video on USB. Engineering will verify the route before transferring the file."],
        ["Backend owner confirmation", "Decide whether vehicle_only and group_confidence_floor are accepted /detection-event outcomes, and confirm Idempotency-Key handling."],
        ["Real lane access later", "Once mounted at a drive-through, provide a privacy-safe view so engineering can set lane, plate, and menu-board zones."],
        ["Privacy decision on plate hashing", "Choose whether to coordinate a future keyed-HMAC migration. Do not send us API keys, SSH passwords, private keys, or license-token contents."],
        ["Credential cleanup", "Rotate the SSH password that was shared in chat and move to key-based access after testing."],
        ["More labeled footage", "Keep the current private 30-event set for development, then collect representative day/night/weather footage toward the 200-plate acceptance set."],
    ], [3600, 5760], small=True)

    document.add_heading("Go / no-go rule", level=1)
    _table(document, ["Go when", "Do not expand when"], [
        ["The field unit sustains its agreed frame rate without current power/thermal throttling; the camera zone is calibrated; visits survive an outage; false identities are zero; and the labeled trial meets the agreed count and plate targets.", "The camera view is uncalibrated, the Pi is throttling, the service reports repeated camera/inference failures, queue age grows without recovery, any wrong customer-to-vehicle match appears, or the sample is too small to support the percentage claim."],
    ], [4680, 4680], small=True)

    document.add_paragraph()
    _callout(
        document,
        "Recommendation",
        "Hold the long run until cooling and non-cellular file transfer are ready. Then repeat the "
        "supervised canary on one Pi, keep make/model off, and treat 98% as a measured acceptance "
        "test—not a configuration setting.",
        LIGHT_AMBER,
        AMBER,
    )

    path = OUT / "Serv_Pi_Executive_Change_Summary_2026-08-05.docx"
    document.save(path)
    return path


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for path in (build_technical(), build_nontechnical()):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
